from __future__ import annotations

import json
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "交通目标检索与跨镜轨迹回溯平台技术说明书_面试版_v5_逻辑细化.docx"

ACCENT = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_font(paragraph, size=10.8, bold=False, color=BLACK):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def set_style_font(style, size, bold=False, color=BLACK):
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    if fill:
        shade_cell(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=9.6, bold=bold, color=BLACK)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(widths[idx].inches * 1440)))


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.8, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.8)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.8)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run_font(r, size=10.4)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run_font(r, size=10.4)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, fill=TABLE_FILL, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), align=WD_ALIGN_PARAGRAPH.LEFT)
    if widths:
        set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(9.2)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, "FFF8E8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(title + "：")
    set_run_font(r1, size=10.5, bold=True, color=RGBColor(122, 90, 0))
    r2 = p.add_run(body)
    set_run_font(r2, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def read_summary():
    path = ROOT / "output" / "cityflow_results.json"
    if not path.exists():
        return {}
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        return data.get("summary", {})
    except Exception:
        return {}


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    set_style_font(styles["Normal"], 10.8)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    set_style_font(styles["Heading 1"], 16, bold=True, color=BLUE)
    set_style_font(styles["Heading 2"], 13, bold=True, color=BLUE)
    set_style_font(styles["Heading 3"], 11.5, bold=True, color=ACCENT)
    set_style_font(styles["List Bullet"], 10.4)
    set_style_font(styles["List Number"], 10.4)
    return doc


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("交通目标检索与跨镜轨迹回溯平台")
    set_run_font(r, size=22, bold=True, color=RGBColor(11, 37, 69))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("项目技术说明书（算法岗面试版）")
    set_run_font(r, size=15, bold=True, color=MUTED)

    rows = [
        ("项目定位", "端到端多任务交通目标检索、身份关联与跨摄像头时空回溯平台"),
        ("核心算法", "Chinese-CLIP 图文双塔检索、ReID 外观重识别、属性识别、跨镜候选边学习、观测链解码、多任务联合损失"),
        ("平台形态", "FastAPI 后端服务 + Streamlit 前端 + 离线预处理脚本 + Docker Compose 部署"),
        ("数据基础", "AICity22 Track1 MTMC 多摄像头车辆数据，含检测、轨迹、摄像头元数据和关键帧裁剪"),
        ("生成日期", datetime.now().strftime("%Y-%m-%d")),
    ]
    add_table(doc, ["项目", "说明"], rows, [Inches(1.45), Inches(5.85)])

    add_callout(
        doc,
        "面试使用建议",
        "不要只说“用了 CLIP、YOLO 和 ReID”。更好的讲法是：我把任务建模成端到端多任务学习问题，前端感知分支输出目标实例和 tracklet，中间表征分支同时学习图文双塔对齐、ReID 身份嵌入和属性预测，后端图匹配分支学习跨镜候选边得分，最终解码为可解释的观测链和推断段。",
    )
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "目录", 1)
    items = [
        "1．项目定位与业务闭环",
        "2．总体平台架构",
        "3．数据资产与数据流向",
        "4．数据获取、目录结构与训练样本构造",
        "5．模型选型与关键技术取舍",
        "6．离线预处理链路",
        "7．在线检索链路",
        "8．目标确认与跨镜回溯链路",
        "9．关键算法细节",
        "10．端到端训练目标与损失函数",
        "11．算法亮点与创新点",
        "12．接口、前端与部署架构",
        "13．面试可讲重点与追问回答",
        "14．当前边界与后续优化方向",
    ]
    add_numbered(doc, items)
    doc.add_page_break()


def build_document():
    summary = read_summary()
    doc = setup_doc()
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "1．项目定位与业务闭环", 1)
    add_para(
        doc,
        "本项目是一个面向道路交通视频的端到端目标检索与轨迹研判平台。它解决的不是单纯“给一段文字找相似图片”的问题，而是把输入视频、文本查询、目标检测、身份特征、属性预测、跨镜候选边和最终轨迹链放在同一个算法框架中建模：输入是多摄像头视频和自然语言查询，输出是候选目标、确认目标、跨摄像头观测链、推断路径和置信度。"
    )
    add_para(
        doc,
        "从算法岗面试角度，可以把项目概括为一个端到端多任务框架：前端感知分支负责目标检测和单摄跟踪，中间表征分支同时学习 Chinese-CLIP 图文对齐、ReID 身份嵌入和属性分类，后端图匹配分支学习 tracklet 之间的跨镜连接概率，最后通过图搜索或序列解码得到目标观测链。工程部署时仍然可以离线建库和在线检索，但算法设计上要把这些环节视为可联合优化的整体。"
    )
    add_table(
        doc,
        ["问题", "传统做法的不足", "本项目的处理方式"],
        [
            ("文本找车", "只依赖关键词或人工翻看，难以在大量目标中快速定位。", "用图文双塔对比学习把文本和图像映射到同一空间，再结合属性监督提升细粒度检索。"),
            ("候选误检", "相似图片可能不是同一目标。", "用 ReID 身份损失和跨镜边损失约束候选目标，使同一身份更接近、错误连接得分更低。"),
            ("跨镜轨迹", "摄像头覆盖不连续，不能直接生成真实连续轨迹。", "把轨迹回溯建模为 tracklet 图上的边分类和路径解码问题，输出观测段和推断段。"),
            ("面试表达", "容易讲成“模型堆叠”。", "按端到端输入、共享表征、多任务损失、图匹配解码和可解释输出讲。"),
        ],
        [Inches(1.25), Inches(2.55), Inches(3.5)],
    )

    add_heading(doc, "2．总体平台架构", 1)
    add_para(doc, "平台工程上采用“离线预处理 + 在线 API 服务 + 前端研判界面”的分层结构；算法上则可以理解为端到端多任务网络。离线侧负责把视频转成可索引数据，在线侧负责检索、确认、回溯和展示；训练侧则通过检测、属性、图文对齐、ReID 和跨镜边监督共同优化。")
    add_code_block(
        doc,
        "多摄像头视频 + 文本查询 + 标注监督\n"
        "  -> 感知 backbone：目标检测、单摄跟踪、目标 crop\n"
        "  -> 表征学习：Chinese-CLIP 双塔、ReID embedding、属性分类 head\n"
        "  -> 候选图构建：tracklet 节点、跨镜候选边、时空拓扑特征\n"
        "  -> 边评分网络：学习 p(edge_ij = same_id)\n"
        "  -> 路径解码：生成 observation nodes、inference segments、candidate paths\n"
        "  -> 端到端损失：L_det + L_clip + L_reid + L_attr + L_edge + L_path\n"
        "  -> 工程输出：候选目标、确认目标、跨摄像头时间线、证据分数、推断置信度"
    )
    add_table(
        doc,
        ["层级", "核心文件或组件", "职责", "面试讲法"],
        [
            ("表现层", "frontend/pages/search.py、confirm.py、trajectory.py、dashboard.py", "提供查询、确认、轨迹展示和统计看板。", "前端服务算法结果的可解释展示。"),
            ("服务层", "api/main.py、api/routes/search.py、backtrack.py", "提供 REST API，封装检索和回溯流程。", "把算法链路产品化为稳定接口。"),
            ("业务层", "src/retrieval、src/stitching、src/backtrack", "查询解析、候选重排、跨镜评分和观测链构建。", "端到端训练时对应双塔对齐、ReID、边分类和路径解码。"),
            ("数据层", "output/cityflow_results.json、clip_vectors.faiss、configs/*.yaml", "保存结构化检测、轨迹、向量和摄像头元数据。", "离线结构化是在线秒级查询的前提。"),
            ("部署层", "Dockerfile.backend、Dockerfile.frontend、docker-compose.yml", "后端、前端、Qdrant 服务编排。", "便于 GPU 推理服务和前端独立部署。"),
        ],
        [Inches(0.9), Inches(1.7), Inches(2.35), Inches(2.35)],
    )

    add_heading(doc, "3．数据资产与数据流向", 1)
    total_dets = summary.get("total_detections", "报告记录为 68,349")
    total_tracks = summary.get("total_tracks", "报告记录为 2,070")
    add_para(
        doc,
        f"项目数据围绕目标实例展开。一个目标实例通常包含 target_id、camera_id、scene_id、frame_id、timestamp、bbox、confidence、attributes、keyframe_path 和可选向量特征。当前报告记录的数据规模为 46 个摄像头、230 个车辆 ID、约 {total_dets} 条检测记录、约 {total_tracks} 条轨迹。"
    )
    add_table(
        doc,
        ["数据对象", "关键字段", "作用"],
        [
            ("TargetInstance", "instance_id、camera_id、timestamp、bbox、attributes、keyframe_path", "一次目标观测，是检索候选和回溯锚点的基础。"),
            ("Tracklet", "tracklet_id、camera_id、start_time、end_time、instances、attributes、avg_reid_vector", "同一摄像头内的连续观测段，是跨镜拼接的图节点。"),
            ("CrossCameraEdge", "source_tracklet_id、target_tracklet_id、score、各分项得分、reasoning", "两个 tracklet 是否属于同一目标的候选连接。"),
            ("TrajectoryResult", "observation_nodes、observation_segments、inference_segments、candidate_paths、evidence", "最终面向用户和前端的轨迹回溯结果。"),
        ],
        [Inches(1.45), Inches(2.6), Inches(3.25)],
    )
    add_callout(
        doc,
        "核心理解",
        "视频不会直接进入在线检索。真正被在线服务使用的是结构化后的目标实例库、轨迹库、图片裁剪和向量索引。这样才能把一次查询压缩到属性匹配和向量检索，而不是临时扫视频。",
    )
    add_heading(doc, "3.1 数据由什么变成什么", 2)
    add_table(
        doc,
        ["阶段", "输入数据", "转换方法", "输出数据", "为什么重要"],
        [
            ("原始采集层", "多摄像头 vdo.avi、MOT 标注 gt.txt、摄像头时间戳。", "按场景和摄像头解析视频帧、目标框、帧号和时间偏移。", "可定位到某帧某目标的原始观测。", "把非结构化视频变成可计算的帧级证据。"),
            ("检测实例层", "帧图像、bbox、vehicle_id、camera_id。", "裁剪目标图像，写入 target_id、frame_id、timestamp、bbox、confidence。", "Detection / TargetInstance。", "这是图文检索和目标确认的最小单位。"),
            ("单摄轨迹层", "同一 camera_id 下同一 vehicle_id 的多次检测。", "按时间排序并聚合关键帧、属性、起止帧和检测数量。", "Tracklet / track 记录。", "摄像头内轨迹是真实观测段，可信度最高。"),
            ("属性与质量层", "目标 crop、标注属性、检测框。", "颜色映射、KMeans 主色分析、车型映射、置信度和质量评分。", "颜色、车型、尺寸、quality_score。", "给检索粗筛和跨镜属性一致性提供可解释证据。"),
            ("共享表征层", "目标 crop 图像、文本描述、身份 ID、属性标签。", "ReID 分支提取外观向量；Chinese-CLIP 图像塔和文本塔学习跨模态对齐；属性 head 预测颜色和车型。", "reid_vector、clip_image_vector、clip_text_vector、attribute_logits。", "把检索、身份关联和属性识别放进统一表征空间。"),
            ("候选图层", "tracklet、camera_id、timestamp、拓扑邻接、ReID 相似度、属性一致性。", "构造 tracklet 图，候选边由时间窗口、空间可达和外观相似度初筛生成。", "edge_features、edge_label、candidate_graph。", "把跨镜轨迹问题转成图上的边分类和路径选择。"),
            ("训练监督层", "文本－图像配对、车辆 ID、属性标签、真实跨镜连接、真实轨迹链。", "分别计算对比损失、ReID 损失、属性损失、边分类损失和路径排序损失。", "L_clip、L_reid、L_attr、L_edge、L_path。", "端到端优化所有关键算法目标。"),
            ("推理索引层", "所有目标实例的向量和元数据。", "向量归一化后写入 FAISS 或 Qdrant，并保留 target_id 到向量行号的映射。", "clip_vectors.faiss、Qdrant collection、cityflow_results.json。", "部署时不扫视频，只查索引和 JSON 元数据。"),
            ("在线结果层", "query_text 或 plate_number、确认的 instance_id。", "双塔向量召回、属性和 ReID 融合排序、跨镜边评分、路径解码。", "候选目标、确认目标、camera_sequence、观测段、推断段、证据分数。", "形成面向业务研判的最终输出。"),
        ],
        [Inches(0.9), Inches(1.35), Inches(1.55), Inches(1.45), Inches(2.05)],
    )

    add_heading(doc, "4．数据获取、目录结构与训练样本构造", 1)
    add_para(doc, "项目的数据不是手工造的，而是围绕 AICity22 Track1 MTMC Tracking 数据集整理。下载脚本指向 HuggingFace 数据集 `AICity2022-Challenge/AICity2022-Track1-MTMC`，本地目录为 `cityflow/AICity22_Track1_MTMC_Tracking`。这个数据集适合本项目，因为它天然包含多摄像头、多场景、车辆 ID、时间戳和 MOT 格式标注，能同时支撑检测、单摄跟踪、ReID、跨镜轨迹和时空约束建模。")
    add_table(
        doc,
        ["数据来源", "项目中对应文件", "字段或内容", "用途"],
        [
            ("AICity22 Track1 MTMC", "cityflow/AICity22_Track1_MTMC_Tracking", "6 个场景、46 个摄像头、train / validation / test 划分。", "多摄像头交通目标检索和跨镜回溯主数据。"),
            ("视频文件", "{split}/{scene}/{camera}/vdo.avi", "每个摄像头一个道路监控视频。", "抽帧、裁剪目标、可视化关键帧。"),
            ("GT 标注", "{split}/{scene}/{camera}/gt/gt.txt", "MOT 格式：frame、id、left、top、width、height、conf、class、visibility。", "生成 detection、tracklet、ReID 正负样本和跨镜真值。"),
            ("检测文件", "{split}/{scene}/{camera}/det/det.txt", "检测候选框。", "可用于对比 YOLO 检测或导入 baseline detection。"),
            ("摄像头时间", "cam_timestamp/{scene}.txt", "每个摄像头的起始时间偏移。", "把 frame_id 换算成真实 timestamp，用于跨镜时间可达性。"),
            ("帧数文件", "cam_framenum/{scene}.txt", "每个摄像头视频帧数。", "验证视频完整性和抽帧范围。"),
            ("摄像头元数据", "configs/cityflow_camera_metadata.yaml", "camera_id、场景、坐标、车道方向、邻接关系。", "地图展示、空间可达性、方向一致性和候选路径生成。"),
        ],
        [Inches(1.25), Inches(1.85), Inches(2.35), Inches(1.85)],
    )
    add_heading(doc, "4.1 目录结构和处理顺序", 2)
    add_code_block(
        doc,
        "cityflow/AICity22_Track1_MTMC_Tracking/\n"
        "  train/S01/c001/vdo.avi, gt/gt.txt, det/det.txt\n"
        "  validation/S02/c006/vdo.avi, gt/gt.txt, det/det.txt\n"
        "  test/S06/c041/vdo.avi, det/det.txt\n"
        "  cam_timestamp/S01.txt ... S06.txt\n"
        "  cam_framenum/S01.txt ... S06.txt\n\n"
        "处理顺序：下载数据 -> 验证视频和标注 -> 解析 MOT 标注 -> 抽取关键帧 crop -> 补全属性 -> 提取 ReID / CLIP 特征 -> 构建索引 -> 启动 API 和前端"
    )
    add_heading(doc, "4.2 训练样本如何从数据中构造", 2)
    add_table(
        doc,
        ["训练任务", "正样本", "负样本", "样本构造细节"],
        [
            ("图文双塔检索", "同一目标 crop 和其中文描述，例如“黑色中型轿车”。", "不同颜色、不同车型或不同车辆 ID 的 crop。", "描述由颜色、车型、尺寸和场景信息生成；hard negative 优先选颜色相同但车型不同、车型相同但颜色不同的目标。"),
            ("ReID 身份学习", "同一 vehicle_id 在不同摄像头或不同帧的 crop。", "不同 vehicle_id 的 crop，优先选择外观相近车辆。", "anchor / positive / negative 从 MOT id 和 camera_id 组合中抽样，支持 Batch Hard Triplet。"),
            ("属性识别", "带颜色、车型、背包、服饰标签的目标 crop。", "非对应类别样本。", "车辆颜色和车型来自标注映射、仿真属性表或确定性补全；类别不均衡时用 Focal Loss。"),
            ("跨镜候选边", "同一 vehicle_id 在不同 camera_id 的相邻 tracklet。", "时间相近但 vehicle_id 不同，或拓扑可达但身份不同的 tracklet 对。", "边特征包括 ReID 相似度、CLIP 相似度、属性一致性、车牌一致性、时间差、距离、方向和拓扑。"),
            ("轨迹路径排序", "同一 vehicle_id 的真实 camera_sequence。", "替换其中一个节点或边后的错误路径。", "训练 path ranking loss，让真实路径得分高于错误路径。"),
        ],
        [Inches(1.15), Inches(1.75), Inches(1.75), Inches(2.65)],
    )
    add_callout(
        doc,
        "像做过项目的细节",
        "样本构造不能只说“正负样本”。要说清楚正样本来自同一 vehicle_id 的跨摄像头观测，负样本不是随机取，而是优先取时间相近、外观相似、颜色或车型相同的 hard negative，因为这些才是真实系统最容易误关联的样本。",
    )

    add_heading(doc, "4.3 可扩展数据来源", 2)
    add_para(doc, "AICity22 是本项目主数据源，适合验证多摄像头车辆轨迹和跨镜回溯。但如果要把项目讲成可持续优化的平台，还需要说明后续如何补充数据。合理做法是：主任务用 AICity22 建立闭环，ReID 分支用车辆 ReID 数据集增强身份表征，中文图文分支用中文图文数据或项目内自动描述做弱监督微调，时空分支接入道路拓扑数据增强可达性判断。")
    add_table(
        doc,
        ["数据类型", "可用来源", "用于哪个模块", "引入方式"],
        [
            ("多摄像头车辆轨迹", "AICity22 Track1 MTMC", "检测、轨迹、跨镜边、路径监督。", "解析 vdo.avi、gt.txt、cam_timestamp、cam_loc，构造 detection、tracklet、camera sequence。"),
            ("车辆 ReID", "VeRi-776、VehicleID、VERI-Wild", "ReID backbone 预训练或微调。", "同 ID 为正样本，不同 ID 为负样本，使用 ID 分类和 Triplet Loss。"),
            ("行人 ReID", "Market-1501、DukeMTMC、MSMT17", "如果平台扩展到行人检索。", "替换或新增行人 ReID 分支，避免用车辆外观模型处理行人。"),
            ("中文图文对", "MUGE、COCO-CN、Noah-Wukong 或项目自动生成描述。", "Chinese-CLIP 双塔微调。", "构造“图像 crop－中文描述”配对，并加入 hard negative。"),
            ("道路拓扑", "OpenStreetMap、摄像头位置和道路连通关系。", "跨镜候选边、路径搜索和不可能路径过滤。", "把摄像头节点映射到路网，计算距离、方向、可达时间窗口。"),
            ("人工反馈", "前端确认、驳回和修正记录。", "重排、边评分和错误案例挖掘。", "把用户确认结果转成正负样本，做周期性再训练。"),
        ],
        [Inches(1.35), Inches(1.65), Inches(1.9), Inches(2.4)],
    )

    add_heading(doc, "5．模型选型与关键技术取舍", 1)
    add_para(doc, "模型选型遵循两个原则：一是先满足交通视频场景的业务约束，二是把高成本模型放在候选规模变小之后使用。检测用 YOLOv8x 是为了保证车辆、行人和非机动车的基础召回；图文检索用 Chinese-CLIP 是因为用户输入是中文自然语言；跨镜身份关联用 ReID 是因为 CLIP 不能承担强身份判别；向量检索用 FAISS / Qdrant 是为了让在线查询不扫描视频。")
    add_table(
        doc,
        ["模块", "选择", "为什么选它", "没有选其他方案的原因"],
        [
            ("目标检测", "YOLOv8x", "精度高、工程成熟、ultralytics 接口稳定，COCO 类别可映射到 person / car / bus / truck / bicycle / motorcycle。", "YOLOv8s 更快但召回和小目标能力较弱；DETR 系列训练和部署成本更高。"),
            ("单摄跟踪", "ByteTrack 思路 / MOT id 聚合", "交通视频里同一摄像头内目标运动连续，按检测框和时间聚合能生成稳定 tracklet。", "DeepSORT 依赖 ReID，前期会增加复杂度；先把单摄轨迹闭环跑通更重要。"),
            ("图文检索", "Chinese-CLIP ViT-L/14 或 ViT-B/16", "中文文本和目标图像进入同一向量空间，图像向量可离线预计算，适合大规模检索。", "原版 CLIP 中文支持弱；BLIP 逐对匹配更慢，更适合 Top-K 精排而非全库召回。"),
            ("ReID", "OSNet / ResNet50 backbone，512 维 embedding", "用于跨摄像头身份一致性，和 CLIP 的语义匹配互补。", "只用 CLIP 容易把语义相似目标误当成同一目标；只用 ReID 又不能支持自然语言查询。"),
            ("属性识别", "颜色 KMeans + 标注映射 + 车型映射", "颜色和车型是交警检索最常用字段，也能解释结果。", "直接上复杂属性网络需要额外标注；当前先做稳定、可解释的属性补全。"),
            ("向量索引", "FAISS IndexFlatIP + Qdrant 预留", "IndexFlatIP 对当前规模精确、实现简单；Qdrant 适合后续服务化和增量索引。", "HNSW / IVF 适合十万级以上，当前阶段没必要先引入近似误差。"),
            ("跨镜匹配", "可学习边评分 + 交通约束", "ReID、属性、车牌、时间、空间、方向一起判断，避免单一视觉特征误关联。", "纯 ReID 缺少道路可达性；纯规则又难以自适应不同场景。"),
        ],
        [Inches(1.0), Inches(1.3), Inches(2.55), Inches(2.45)],
    )
    add_heading(doc, "5.1 模型输入输出细节", 2)
    add_table(
        doc,
        ["模型", "输入", "输出", "关键处理"],
        [
            ("YOLOv8x", "BGR 视频帧，输入尺寸 640 x 640。", "bbox、class_id、confidence。", "conf=0.5，NMS IoU=0.45；COCO class 0->pedestrian，1/3->non_motor_vehicle，2/5/7->vehicle。"),
            ("Chinese-CLIP 图像塔", "目标 crop RGB 图像。", "归一化 image embedding。", "离线批量编码，保存到 JSON / NPY / FAISS / Qdrant。"),
            ("Chinese-CLIP 文本塔", "用户 query_text 或生成描述。", "归一化 text embedding。", "在线编码查询；注意 Chinese-CLIP 要使用 ChineseCLIPModel 专用类，不能简单用通用 CLIPModel。"),
            ("ReID 分支", "resize 后的目标 crop，常见尺寸 256 x 128。", "512 维 L2 归一化外观 embedding。", "tracklet 内多帧向量可均值聚合为 avg_reid_vector。"),
            ("边评分 MLP", "两个 tracklet 的相似度和时空特征。", "p(edge=same_id)。", "输入包含 S_reid、S_clip、S_attr、S_plate、Δt、distance、direction、topology。"),
        ],
        [Inches(1.2), Inches(1.65), Inches(1.55), Inches(2.9)],
    )
    add_heading(doc, "5.2 项目里容易被问到的技术取舍", 2)
    add_bullets(doc, [
        "为什么要双塔：图库图像向量可以离线预计算，在线只算一次文本向量，适合实时检索。",
        "为什么还要 ReID：CLIP 解决语义匹配，不解决强身份判别；跨镜同一车辆判断必须引入 ReID 或车牌等身份证据。",
        "为什么先属性粗筛：颜色和车型是强业务字段，能减少向量误召回，也让候选排序更可解释。",
        "为什么区分 observation 和 inference：摄像头视野不连续，摄像头之间没有真实画面，只能给推断段和置信度。",
        "为什么保留 FAISS 和 Qdrant 两条路径：FAISS 便于本地精确检索和实验复现，Qdrant 适合服务化部署和后续增量数据。"
    ])

    add_heading(doc, "5.3 模型验证与关键参数", 2)
    add_para(doc, "项目里模型不是只在文档中列名，而是需要在上线前做最小可用验证：检测模型要能在真实帧上输出车辆框，Chinese-CLIP 要能对中文文本和目标 crop 输出非零归一化向量，ReID 分支要能对同一车辆多帧输出稳定 embedding，索引端要能保证向量维度、metadata 数量和候选图片一一对应。")
    add_table(
        doc,
        ["参数或检查项", "当前取值或做法", "作用"],
        [
            ("YOLO 输入尺寸", "640 x 640", "兼顾速度和车辆小目标检测效果。"),
            ("检测置信度阈值", "conf = 0.5", "过滤明显低质量检测框，降低后续裁剪噪声。"),
            ("NMS 阈值", "IoU = 0.45", "去除同一车辆的重复框。"),
            ("ByteTrack / 单摄跟踪", "max_age = 30、min_hits = 3、iou = 0.3", "保证短时遮挡时轨迹不断裂，同时避免短暂误检成为稳定轨迹。"),
            ("ReID 输入", "256 x 128 crop", "车辆和行人 ReID 常用长宽比例，便于 backbone 学习外观结构。"),
            ("ReID 输出", "512 维 L2 归一化 embedding", "用于跨帧、跨镜外观相似度。"),
            ("CLIP batch size", "16", "图像塔离线批处理，显存占用可控。"),
            ("ReID batch size", "32", "离线提取外观特征时提升吞吐。"),
            ("FAISS 召回规模", "Top-200", "先扩大语义召回，再通过属性和 track 去重收敛到前端 Top-K。"),
            ("最终融合权重", "0.6 * clip_score + 0.4 * attr_score", "让语义匹配主导排序，同时保留颜色、车型等强约束。"),
            ("跨镜速度约束", "20-80 km/h", "过滤不符合城市道路速度范围的候选连接。"),
            ("回溯最小置信度", "min_confidence = 0.5", "低置信跨镜边不直接作为可靠轨迹输出。"),
        ],
        [Inches(1.7), Inches(2.25), Inches(3.35)],
    )

    add_heading(doc, "5.4 Chinese-CLIP 使用的 embedding 模型", 2)
    add_para(doc, "这里说的 embedding 模型，不是单独再接一个普通文本向量模型，而是 Chinese-CLIP 自身的双塔 embedding。Chinese-CLIP 有两个编码器：图像编码器把车辆 crop 编成 image embedding，文本编码器把中文查询编成 text embedding。两个 embedding 会被投影到同一个语义向量空间，经过 L2 归一化后，可以直接用内积或余弦相似度比较。")
    add_table(
        doc,
        ["模块", "可选模型", "输入", "输出 embedding", "在项目中的用法"],
        [
            ("图像 embedding 模型", "OFA-Sys/chinese-clip-vit-large-patch14 或 cn_clip 的 ViT-B-16 / ViT-L-14", "车辆、行人或非机动车目标 crop。", "image embedding，常见维度为 512 或 768，具体取决于模型版本。", "离线批量提取并写入 FAISS / Qdrant，作为图库向量。"),
            ("文本 embedding 模型", "与图像塔配套的 Chinese-CLIP text encoder", "用户中文描述，例如“蓝色轿车”“白色 SUV”。", "text embedding，维度必须和 image embedding 完全一致。", "在线查询时实时提取，用于向量召回。"),
            ("投影与归一化", "Chinese-CLIP projection head + L2 normalization", "图像塔和文本塔的原始特征。", "同一向量空间中的单位向量。", "用内积近似余弦相似度，方便向量检索和分数融合。"),
            ("向量索引", "FAISS IndexFlatIP 或 Qdrant cosine / dot product", "预计算 image embedding 和在线 text embedding。", "Top-K 相似目标列表。", "把“中文描述”变成候选车辆 crop。"),
        ],
        [Inches(1.25), Inches(1.85), Inches(1.45), Inches(1.4), Inches(1.35)],
    )
    add_callout(
        doc,
        "面试中要讲清楚",
        "如果面试官问“embedding 模型是什么”，可以回答：项目里 embedding 主要来自 Chinese-CLIP 双塔模型，不是普通文本 embedding。图像塔和文本塔分别产生 image embedding 与 text embedding，并通过对比学习对齐到同一个语义空间。检索时图库 image embedding 离线生成，用户 query 的 text embedding 在线生成，两者做余弦相似度得到候选目标。"
    )
    add_table(
        doc,
        ["容易混淆的问题", "正确说法"],
        [
            ("是不是只用了文本 embedding？", "不是。文搜图必须同时有 text embedding 和 image embedding，二者在同一空间比较。"),
            ("能不能换成 BGE、text2vec 这类中文文本向量？", "不能直接替代文搜图，因为普通文本 embedding 只编码文本，不能和图像 crop 直接比较；除非再训练一个图像－文本对齐模型。"),
            ("为什么同一个项目会出现 512 维和 768 维？", "不同 Chinese-CLIP backbone 的输出维度不同。建库和查询必须固定同一个模型版本，不能用 768 维图库配 512 维查询。"),
            ("为什么要归一化？", "归一化后向量长度不影响相似度，内积就可以近似余弦相似度，便于 FAISS / Qdrant 检索。"),
        ],
        [Inches(2.25), Inches(5.05)],
    )

    add_heading(doc, "6．离线预处理链路", 1)
    add_para(doc, "离线预处理的目标是把不可直接查询的视频文件转成结构化、可索引、可展示的数据资产。项目中主要由 scripts/cityflow_all_tracks.py 和 scripts/clip_feature_pipeline.py 完成。")
    add_code_block(
        doc,
        "AICity22 视频与 MOT 标注\n"
        "  -> 解析 gt.txt 和 cam_timestamp\n"
        "  -> 提取每辆车关键帧并裁剪目标图像\n"
        "  -> 生成 detection 和 track 记录\n"
        "  -> 分配或补全颜色、车型等属性\n"
        "  -> 对裁剪图像做颜色分析与描述生成\n"
        "  -> ReID 图像外观向量编码，用于跨镜身份一致性\n"
        "  -> Chinese-CLIP 图像塔编码，用于图文检索\n"
        "  -> 写入 cityflow_results.json、reid_vectors、clip_vectors.npy / faiss、Qdrant（可选）"
    )
    add_table(
        doc,
        ["步骤", "输入", "处理方法", "输出"],
        [
            ("检测和轨迹读取", "vdo.avi、gt.txt、cam_timestamp", "解析 MOT 格式标注，按车辆 ID 和摄像头聚合。", "detections、tracks。"),
            ("关键帧裁剪", "视频帧、bbox", "按车辆轨迹均匀选择关键帧，保存 crop 图片。", "cityflow_crops 下的目标图片。"),
            ("属性补全", "vehicle_id、标注或映射表", "颜色和车型采用标注映射或确定性分配，保证同一 vehicle_id 属性稳定。", "attributes 字段。"),
            ("颜色细化", "目标 crop", "KMeans 主色聚类，结合 GT color_id 做细粒度颜色名。", "color_refined、dominant_color_rgb。"),
            ("ReID 编码", "目标 crop 或同一 tracklet 的多帧 crop", "使用 ReID backbone 提取 512 维外观 embedding，tracklet 级可做均值聚合。", "reid_vector、avg_reid_vector。"),
            ("双塔编码", "目标 crop、中文描述或在线 query_text", "Chinese-CLIP 图像塔离线编码图像，文本塔在线编码查询文本，二者在共享嵌入空间中计算相似度。", "clip_image_vector、clip_text_vector、FAISS 或 Qdrant 索引。"),
        ],
        [Inches(1.25), Inches(1.55), Inches(2.65), Inches(1.85)],
    )

    add_heading(doc, "6.1 关键预处理细节", 2)
    add_table(
        doc,
        ["处理点", "具体做法", "为什么重要"],
        [
            ("MOT 标注解析", "从 gt.txt 读取 frame、id、left、top、width、height、confidence、class、visibility 等字段；按 camera_id 和 vehicle_id 组织为 detection 列表。", "后续检索、ReID 样本构造和跨镜轨迹都依赖同一个实例索引，标注解析错误会导致全链路错位。"),
            ("时间戳对齐", "读取 cam_timestamp / cam_framenum 等元数据，将 frame_id 转换成真实 timestamp；如果不同摄像头 fps 不一致，需要按摄像头单独换算。", "跨镜回溯不是只看帧号，而是比较不同摄像头之间的真实时间先后。"),
            ("关键帧抽取", "对同一 vehicle_id 在同一 camera_id 下的轨迹片段按时间均匀抽取若干帧，避免只取首帧导致遮挡、模糊或角度单一。", "提高 CLIP、ReID 和属性识别的鲁棒性，也让前端候选图更有代表性。"),
            ("目标裁剪", "优先使用 bbox 从视频帧中裁剪目标；若输入是整帧图且 bbox 为全图，则先用 YOLOv8x 检测车辆区域再裁剪。", "CLIP 和 ReID 都应看目标本身，而不是被道路、背景和其他车辆干扰。"),
            ("实例 ID 设计", "target_id 使用类似 CF3_c001_V0034_000001 的格式，track_id 使用摄像头和车辆 ID 组合；回溯时从 target_id 中解析 V0034 作为 vehicle_id。", "保证前端候选、后端检索、跨镜聚合和轨迹接口能通过同一 ID 体系串起来。"),
            ("属性补全", "颜色、车型优先使用标注或映射表；缺失时用确定性规则按 vehicle_id 补全，避免同一车辆在不同帧属性随机跳变。", "属性是粗筛和解释字段，稳定性比花哨模型更重要。"),
            ("颜色细化", "对 crop 做 KMeans 主色聚类，结合颜色中心表和 GT color_id 得到 color_refined 与 dominant_color_rgb。", "面试时可以说明不是只靠中文词典，而是把图像统计结果也写入结构化字段。"),
            ("向量归一化", "CLIP 文本向量、CLIP 图像向量和 ReID embedding 均做 L2 normalization，FAISS IndexFlatIP 的内积近似余弦相似度。", "归一化后分数尺度更稳定，便于融合排序和阈值设置。"),
            ("索引落盘", "结构化结果写 cityflow_results.json，向量写 FAISS / Qdrant；服务启动时加载 metadata 与向量索引。", "把耗时的视觉编码放到离线阶段，在线查询只做文本编码和向量检索。"),
        ],
        [Inches(1.35), Inches(3.55), Inches(2.4)],
    )

    add_heading(doc, "6.2 项目中容易踩坑的处理", 2)
    add_table(
        doc,
        ["问题", "现象", "处理方式"],
        [
            ("Chinese-CLIP 模型接口差异", "cn_clip 与 transformers 的 ChineseCLIPModel 加载方式、输出维度和方法名可能不同。", "建库和查询必须使用同一模型、同一维度、同一归一化方式；索引文件需要记录 feature_dim 和 model_name。"),
            ("CLIP 维度不一致", "有的配置写 CN-CLIP-ViT-L-14 为 768 维，而部分在线接口使用 ViT-B-16 为 512 维。", "不能混用索引；生成索引前固定模型版本，服务加载时校验 FAISS 维度与文本向量维度。"),
            ("整帧图直接入库", "如果把整张道路画面编码为目标图像，检索会被背景和周围车辆影响。", "检测到 bbox 覆盖全图时，先做 YOLO 裁剪；无法裁剪时才作为低置信 fallback。"),
            ("unknown 属性误删", "颜色或车型识别缺失时，如果强过滤会把正确目标删掉。", "属性粗筛只排除明确冲突项，对 unknown 保留并降低属性分。"),
            ("FAISS 行号与 metadata 错位", "检索返回的向量行号映射到错误 detection，前端候选图和分数不一致。", "向量写入、metadata 写入和 cityflow_results.json 更新必须在同一排序下完成，并做数量校验。"),
            ("跨镜时间顺序反常", "不同摄像头帧号不能直接比较，可能出现先后顺序错误。", "统一转 timestamp，再按摄像头起始时间和 fps 校准。"),
            ("同一车辆多帧刷屏", "Top-K 全是同一个 track 的相邻帧。", "在线返回前按 track_id 聚合去重，只保留该 track 的最佳候选。"),
            ("规则回溯像假结果", "如果只按 vehicle_id 聚合，候选路径缺少置信度解释。", "在端到端方案中引入 ReID、时空、方向和拓扑边评分，让每条跨镜连接都有 score 组成。"),
        ],
        [Inches(1.6), Inches(2.55), Inches(3.15)],
    )

    add_heading(doc, "7．在线检索链路", 1)
    add_para(doc, "在线检索接口位于 /api/v1/search/query。它的核心不是一次性用 CLIP 全库搜索，而是先从文本中提取结构化条件，再用属性粗筛降低候选规模，最后用向量相似度和属性分做融合排序。")
    add_code_block(
        doc,
        "用户输入：蓝色轿车\n"
        "  -> 查询解析：color=蓝色，vehicle_type=轿车\n"
        "  -> 属性粗筛：过滤颜色和车型明显不匹配的 detection\n"
        "  -> Chinese-CLIP 文本塔编码：生成归一化查询向量\n"
        "  -> FAISS IndexFlatIP：返回 Top-200 语义相似检测\n"
        "  -> 融合打分：final_score = 0.6 * clip_score + 0.4 * attr_score\n"
        "  -> 按 track_id 聚合去重\n"
        "  -> 返回 Top-K 候选：图片、属性、分数、轨迹摘要"
    )
    add_table(
        doc,
        ["环节", "算法点", "为什么这样做"],
        [
            ("查询解析", "规则词典、口语化映射、车牌正则、颜色和车型提取。", "把自然语言转为结构化条件，提高可控性和可解释性。"),
            ("属性粗筛", "按 color 和 vehicle_type 过滤，unknown 不强制排除。", "减少向量排序空间，并避免属性缺失导致误删。"),
            ("双塔精排", "图像塔特征离线预计算，文本塔在线编码查询，归一化后用内积近似余弦相似度。", "补足关键词无法表达的视觉语义，并保持检索效率。"),
            ("融合打分", "0.6 CLIP 相似度 + 0.4 属性得分。", "兼顾语义匹配和业务属性匹配。"),
            ("轨迹聚合", "同一 track_id 只返回一个代表候选。", "避免同一车辆连续多帧刷屏，提高候选多样性。"),
        ],
        [Inches(1.25), Inches(2.8), Inches(3.25)],
    )
    add_heading(doc, "7.1 在线接口中的数据对象变化", 2)
    add_table(
        doc,
        ["内部对象", "由什么生成", "包含什么", "下一步去向"],
        [
            ("query_features", "query_text 经过关键词、颜色、车型、车牌等解析规则。", "color、vehicle_type、plate、direction、time_range 等结构化条件。", "送入属性粗筛和结果解释。"),
            ("text_embedding", "Chinese-CLIP 文本塔对 query_text 编码。", "L2 归一化后的文本向量，维度必须与图像索引一致。", "送入 FAISS / Qdrant 做向量召回。"),
            ("filtered_detections", "cityflow_results.json 中的 detection 记录经过属性规则过滤。", "保留不冲突或属性未知的候选实例。", "减少向量排序和融合打分的候选规模。"),
            ("clip_scores", "text_embedding 与 clip_image_vector 做内积相似度。", "每个候选的语义匹配分数。", "与属性分合成 final_score。"),
            ("attr_score", "query_features 与 detection.attributes 比较。", "颜色、车型、车牌等字段匹配度。", "参与 final_score = 0.6 * clip_score + 0.4 * attr_score。"),
            ("scored_candidates", "向量分与属性分融合。", "instance_id、track_id、camera_id、timestamp、keyframe、scores、attributes。", "按 final_score 排序。"),
            ("dedup_results", "scored_candidates 按 track_id 聚合。", "每条轨迹只保留最佳代表帧，并附带轨迹摘要。", "返回给前端候选确认界面。"),
        ],
        [Inches(1.45), Inches(1.9), Inches(2.45), Inches(1.5)],
    )
    add_heading(doc, "7.2 请求和响应字段", 2)
    add_table(
        doc,
        ["字段", "方向", "说明"],
        [
            ("query_text", "请求", "用户自然语言，例如“蓝色轿车”“白色 SUV 从路口右转”。"),
            ("top_k", "请求", "返回候选数量，配置中默认检索 top_k 可设为 20。"),
            ("instance_id", "响应", "单个候选实例 ID，用于前端确认和后续回溯。"),
            ("track_id", "响应", "单摄轨迹 ID，用于聚合相邻帧并展示轨迹摘要。"),
            ("camera_id", "响应", "候选出现的摄像头编号。"),
            ("timestamp", "响应", "候选关键帧时间，用于时间轴和跨镜排序。"),
            ("attributes", "响应", "颜色、车型、细化颜色等可解释属性。"),
            ("scores", "响应", "clip_score、attr_score、final_score 等分数组成。"),
            ("keyframe", "响应", "前端展示用目标截图路径。"),
        ],
        [Inches(1.35), Inches(0.9), Inches(5.05)],
    )
    add_callout(
        doc,
        "面试细节",
        "如果被问为什么不用纯 CLIP，回答可以是：纯 CLIP 对颜色、车型等细粒度属性可能不稳定，而且全量向量搜索会返回视觉语义相似但业务属性不合的结果；属性粗筛能降低搜索空间，CLIP 精排则保留语义泛化能力。",
    )

    add_heading(doc, "8．目标确认与跨镜回溯链路", 1)
    add_para(doc, "用户确认候选后，系统从“文本描述”切换到“确定目标实例”。这是检索系统向轨迹研判系统转变的关键，因为后续不再仅依赖描述，而是围绕 instance_id、track_id 或 vehicle_id 查找同目标在多摄像头中的观测记录。")
    add_code_block(
        doc,
        "用户确认 candidate：CF3_c001_V0034_000001\n"
        "  -> 解析 vehicle_id：V0034\n"
        "  -> 查询 vehicle_id -> detections 倒排索引\n"
        "  -> 按 camera_id 分组\n"
        "  -> 每个摄像头内按 timestamp 排序\n"
        "  -> 按首次出现时间构建 camera_sequence\n"
        "  -> 生成每个摄像头的到达时间、离开时间、关键帧和检测数\n"
        "  -> 构建相邻摄像头之间的推断段和置信度\n"
        "  -> 返回完整 trajectory 结构"
    )
    add_table(
        doc,
        ["输出内容", "含义", "展示方式"],
        [
            ("camera_sequence", "目标按时间经过的摄像头列表。", "轨迹页时序表和地图连线。"),
            ("observation_nodes", "目标在某摄像头被实际看到的节点。", "点位、关键帧和时间。"),
            ("observation_segments", "摄像头视野内的真实观测轨迹。", "实线或摄像头内片段说明。"),
            ("inference_segments", "两个摄像头之间无连续画面的推断连接。", "虚线、时间差和置信度。"),
            ("evidence", "外观、属性、时间、空间、车牌等证据。", "研判摘要和置信度解释。"),
        ],
        [Inches(1.55), Inches(3.0), Inches(2.75)],
    )

    add_heading(doc, "8.1 从文搜图到轨迹拼接的触发逻辑", 2)
    add_para(doc, "本项目的主线应该讲成两段式，而不是一上来就做全局轨迹拼接。第一段是文搜图：用户输入中文描述，系统用 Chinese-CLIP 和属性规则返回候选车辆。第二段才是轨迹拼接与路径回溯：用户在候选列表中确认某一辆车后，系统以这个 instance_id / track_id / vehicle_id 作为锚点，向前和向后寻找同一目标在其他摄像头中的观测。")
    add_code_block(
        doc,
        "阶段 1：文搜图检索\n"
        "用户中文 query -> text embedding -> 向量召回 -> 属性融合排序 -> 候选车辆列表\n\n"
        "阶段 2：人工或系统确认目标\n"
        "候选车辆 -> 确认 instance_id / track_id / vehicle_id -> 形成锚点 tracklet\n\n"
        "阶段 3：轨迹拼接\n"
        "锚点 tracklet -> 搜索同摄像头轨迹、重叠摄像头轨迹、无重叠摄像头候选轨迹 -> 计算连接分数\n\n"
        "阶段 4：路径回溯\n"
        "高分连接边 -> 构建观测链 -> 补充推断段 -> 输出 camera_sequence、observation_segments、inference_segments 和置信度"
    )
    add_table(
        doc,
        ["阶段", "输入是什么", "输出变成什么", "关键算法"],
        [
            ("文搜图", "中文描述和离线图像 embedding 库。", "一组候选目标实例，每个候选有图片、属性、分数和 track_id。", "Chinese-CLIP 双塔 embedding、属性粗筛、FAISS / Qdrant 检索。"),
            ("目标确认", "候选目标实例。", "一个确定的锚点目标，例如 CF3_c001_V0034_000001。", "人工确认或 Top-1 自动确认；面试中建议强调人工确认更稳。"),
            ("轨迹拼接", "锚点 tracklet 及其前后时间窗口内的候选 tracklet。", "候选跨镜连接边和边置信度。", "重叠区域用 DTW / 轨迹相似度；无重叠区域用 ReID、属性、时间、拓扑和方向约束。"),
            ("路径回溯", "候选连接边组成的有向图。", "目标完整经过序列、真实观测段和无画面的推断段。", "图搜索、路径打分、置信度校准和证据解释。"),
        ],
        [Inches(1.0), Inches(2.0), Inches(2.25), Inches(2.05)],
    )

    add_heading(doc, "8.2 有覆盖区域与无覆盖区域的轨迹拼接策略", 2)
    add_para(doc, "轨迹拼接不能只说“用 ReID 拼接”。在交通视频中，不同摄像头之间有两种情况：一种是视野存在重叠，车辆可能在两个摄像头中同时或近似同时出现；另一种是没有重叠，车辆从一个摄像头消失后，需要经过一段道路才能到另一个摄像头。两种情况的信息条件不同，所以算法也不同。")
    add_table(
        doc,
        ["场景", "可用信息", "推荐算法", "判断逻辑", "输出"],
        [
            ("有覆盖区域", "两个摄像头视野中存在共同道路区域，车辆轨迹在空间或时间上有重叠。", "DTW 动态时间规整 + 轨迹形状相似度 + ReID 辅助校验。", "把两个摄像头中的轨迹点序列对齐，比较轨迹形状、运动方向、速度变化和时间同步程度。", "判断两个 tracklet 是否其实是同一车辆在重叠区域的连续观测。"),
            ("无覆盖区域", "两个摄像头之间没有连续画面，只知道离开时间、到达时间、摄像头位置和候选目标外观。", "ReID 相似度 + 属性一致性 + 时间可达性 + 路网拓扑 + 方向一致性。", "判断车辆是否能在合理时间内从上游摄像头到达下游摄像头，并且外观和属性是否一致。", "生成 inference_segment，即没有真实画面但算法推断的跨镜连接。"),
            ("不确定区域", "摄像头覆盖关系不清楚，或者缺少准确位置。", "保守候选生成 + 低置信度标记 + 人工复核。", "不强行给出高置信连接，只输出候选路径和证据分。", "避免把推断结果包装成确定轨迹。"),
        ],
        [Inches(1.05), Inches(1.75), Inches(1.75), Inches(1.95), Inches(0.8)],
    )
    add_para(doc, "DTW 的作用是解决两个轨迹序列长度不同、采样频率不同、车辆速度略有变化的问题。比如摄像头 A 中车辆轨迹有 20 个点，摄像头 B 中对应轨迹有 15 个点，不能简单逐点相减；DTW 会寻找一条最小代价对齐路径，让两个序列在时间轴上弹性匹配，再计算整体轨迹差异。")
    add_code_block(
        doc,
        "有覆盖区域连接分数：\n"
        "S_overlap = 0.35 * S_dtw + 0.25 * S_shape + 0.20 * S_reid + 0.10 * S_time_sync + 0.10 * S_direction\n\n"
        "其中：\n"
        "S_dtw 表示两条轨迹点序列经过 DTW 对齐后的相似度；\n"
        "S_shape 表示轨迹整体形状是否一致，例如转弯方向、运动弧线和位移趋势；\n"
        "S_reid 表示目标外观是否相似；\n"
        "S_time_sync 表示两路视频时间是否能对上；\n"
        "S_direction 表示车辆运动方向是否一致。"
    )
    add_code_block(
        doc,
        "无覆盖区域连接分数：\n"
        "S_non_overlap = 0.30 * S_reid + 0.20 * S_temporal + 0.20 * S_topology + 0.15 * S_attribute + 0.10 * S_direction + 0.05 * S_plate\n\n"
        "其中：\n"
        "S_reid 是跨摄像头外观相似度；\n"
        "S_temporal 是离开上游摄像头到进入下游摄像头的时间是否合理；\n"
        "S_topology 是两个摄像头在路网中是否可达；\n"
        "S_attribute 是颜色、车型等属性是否一致；\n"
        "S_direction 是上下游运动方向是否匹配；\n"
        "S_plate 是车牌证据，如果没有车牌则给中性分而不是直接置零。"
    )

    add_heading(doc, "9．关键算法细节", 1)
    add_heading(doc, "9.1 Chinese-CLIP 双塔图文检索", 2)
    add_para(doc, "Chinese-CLIP 属于典型的图文双塔模型：图像塔负责把目标裁剪图编码成视觉向量，文本塔负责把中文查询编码成文本向量。两个塔不需要在检索时逐对交互，而是各自独立编码后在共享向量空间中计算相似度。这一点很关键，因为图像向量可以离线预计算，在线查询只需要计算一次文本向量，再和已有图像向量做近邻检索。")
    add_table(
        doc,
        ["组成", "输入", "输出", "在本项目中的作用"],
        [
            ("图像塔", "车辆或行人 crop 图像。", "归一化图像 embedding。", "离线写入 FAISS / Qdrant，作为图文检索图库。"),
            ("文本塔", "用户输入的中文描述，例如“蓝色轿车”。", "归一化文本 embedding。", "在线生成查询向量，和图像向量计算相似度。"),
            ("相似度层", "文本向量和图像向量。", "余弦相似度或归一化内积。", "返回语义上最匹配的候选目标。"),
            ("元数据层", "target_id、camera_id、timestamp、attributes。", "候选详情和业务字段。", "把向量召回结果还原为可展示、可确认的目标实例。"),
        ],
        [Inches(1.15), Inches(1.8), Inches(1.7), Inches(2.65)],
    )
    add_para(doc, "双塔模型和交叉编码模型的区别也适合面试展开。双塔模型牺牲了一部分逐对细粒度交互能力，但可以预计算图库向量，适合大规模检索；交叉编码模型需要把每个文本－图像对一起输入模型，精度可能更高，但在线成本随候选数线性增长，适合 Top-K 后的精排。")
    add_bullets(doc, [
        "输入：用户自然语言查询和离线保存的目标 crop 图像向量。",
        "处理：文本 tokenization、文本塔编码、向量归一化、FAISS Top-K 检索。",
        "输出：每个 detection 的 clip_score。",
        "优点：能够处理“蓝色轿车”“白色面包车”等自然语言描述，不需要用户记住数据库字段。",
        "边界：CLIP 解决的是“描述和图片是否匹配”，不是“跨摄像头是否同一身份”。因此它不能替代车牌、ReID 和时空约束。"
    ])

    add_heading(doc, "9.2 属性粗筛与融合排序", 2)
    add_para(doc, "属性粗筛的核心是把高维向量检索前置一个低成本、可解释的过滤器。颜色和车型来自离线属性字段，查询解析来自规则词典。对于 unknown 属性，代码中采用不过度惩罚的策略，避免因为属性缺失错杀候选。")
    add_code_block(
        doc,
        "attr_score = confidence\n"
        "if keyword matched:\n"
        "    attr_score = min(0.5 + match_count * 0.15 + confidence * 0.3, 0.99)\n\n"
        "final_score = 0.6 * clip_score + 0.4 * attr_score"
    )

    add_heading(doc, "9.3 ReID 外观重识别特征", 2)
    add_para(doc, "ReID 的目标不是理解自然语言，而是判断不同摄像头中的两个目标外观是否像同一个身份。在本项目中，ReID 是跨镜拼接的核心视觉证据之一。FeatureExtractor 中设计了 ReID 分支，输入目标 crop，输出 512 维 L2 归一化外观向量；同一 tracklet 中多帧目标可以聚合为 avg_reid_vector，用于两个 tracklet 之间的外观相似度计算。")
    add_table(
        doc,
        ["对比项", "Chinese-CLIP 双塔特征", "ReID 特征"],
        [
            ("核心问题", "这段中文描述和这张目标图是否语义匹配？", "这两个摄像头里的目标是否像同一个身份？"),
            ("输入", "中文文本、目标图像。", "目标裁剪图，通常可聚合同一 tracklet 的多帧图像。"),
            ("输出", "图文共享空间中的 embedding。", "身份判别导向的外观 embedding。"),
            ("使用位置", "文本检索候选召回和排序。", "跨镜候选边评分、轨迹拼接、同目标确认。"),
            ("局限", "不擅长做强身份判别。", "不理解自然语言，且受视角、遮挡、光照影响。"),
            ("互补关系", "负责“找像描述的目标”。", "负责“确认跨镜是否同一目标”。"),
        ],
        [Inches(1.1), Inches(3.1), Inches(3.1)],
    )
    add_code_block(
        doc,
        "ReID 侧：crop_image -> ReID backbone -> 512 维 embedding -> L2 normalize\n"
        "Tracklet 聚合：avg_reid_vector = mean(reid_vector_1, ..., reid_vector_n)\n"
        "外观相似度：S_reid = (cos(avg_reid_i, avg_reid_j) + 1) / 2\n"
        "回退策略：如果 ReID 向量缺失，则用 CLIP 图像向量作为备用外观特征"
    )
    add_callout(
        doc,
        "面试表达",
        "可以强调 ReID 和 CLIP 不是重复模块。CLIP 面向图文语义，适合“文字找目标”；ReID 面向身份判别，适合“跨摄像头判断是不是同一辆车或同一个人”。二者都输出向量，但训练目标、相似度含义和使用位置不同。",
    )

    add_heading(doc, "9.4 跨镜候选连接评分", 2)
    add_para(doc, "src/stitching/scoring.py 中的 CrossCameraScorer 是更完整的跨镜拼接评分框架。它把两个单摄 tracklet 之间是否属于同一目标的问题拆成多个可解释分项，然后加权求和并扣除路径分叉和观测缺失惩罚。")
    add_code_block(
        doc,
        "S = w_reid * S_appearance\n"
        "  + w_attr * S_attribute\n"
        "  + w_plate * S_plate\n"
        "  + w_temporal * S_temporal\n"
        "  + w_topology * S_spatial\n"
        "  + 0.10 * S_direction\n"
        "  - penalty_path_divergence\n"
        "  - penalty_observation_missing"
    )
    add_table(
        doc,
        ["分项", "计算逻辑", "面试解释"],
        [
            ("外观相似度", "优先 ReID 向量余弦相似度，缺失时回退 CLIP，映射到 0 到 1。", "解决不同摄像头下同一目标外观一致性。"),
            ("属性一致性", "比较颜色、车型、服饰、背包等属性，缺失给中性分。", "避免强制惩罚缺失字段。"),
            ("车牌一致性", "相同为 1，不同为 0，缺失为 0.5。", "车牌是车辆强身份信息，冲突应强烈否决。"),
            ("时间可达性", "根据 source.end_time 到 target.start_time 的时间差和合理速度区间评分。", "约束不可能的跨镜跳转。"),
            ("空间可达性", "基于摄像头拓扑可达性、距离和估算速度评分。", "把交通路网知识引入视觉匹配。"),
            ("方向一致性", "比较运动方向角度，方向相反降低得分。", "道路方向能过滤逆向或不合理连接。"),
        ],
        [Inches(1.25), Inches(3.0), Inches(3.05)],
    )

    add_heading(doc, "9.5 有覆盖区域：DTW 与轨迹相似度", 2)
    add_para(doc, "当两个摄像头存在视野重叠时，跨镜拼接更像“同一段运动轨迹在两个视角下是否能对齐”。此时不能只看 ReID，因为重叠区域里可能有多辆外观相似车辆同时出现；也不能只看时间，因为两个摄像头可能有帧率差异或时间戳轻微偏移。因此更合理的做法是把轨迹点序列拿出来做 DTW 对齐，再结合轨迹形状相似度判断。")
    add_table(
        doc,
        ["输入", "处理", "输出", "解释"],
        [
            ("tracklet A 的轨迹点序列", "取 bbox 中心点、速度方向、时间戳，形成 [(x, y, t), ...]。", "标准化后的轨迹序列 A。", "表示车辆在摄像头 A 视野中的运动。"),
            ("tracklet B 的轨迹点序列", "同样提取 bbox 中心点和时间戳，并映射到同一平面或相对坐标系。", "标准化后的轨迹序列 B。", "表示候选车辆在摄像头 B 视野中的运动。"),
            ("DTW 对齐", "允许一个点对齐多个点，寻找总代价最小的对齐路径。", "dtw_distance 和对齐路径。", "解决两个轨迹长度不同、采样不同步的问题。"),
            ("轨迹形状相似度", "比较方向变化、速度趋势、起终点关系、转弯形态。", "shape_score。", "判断两条轨迹是否在几何形态上像同一段运动。"),
            ("综合判断", "融合 DTW、形状、ReID、时间同步和方向一致性。", "overlap_edge_score。", "作为重叠摄像头之间的连接置信度。"),
        ],
        [Inches(1.45), Inches(2.35), Inches(1.65), Inches(1.85)],
    )
    add_para(doc, "DTW 的直观解释是：如果两辆车真的是同一辆车，那么即使一个摄像头采了 20 个轨迹点，另一个摄像头只采了 15 个点，它们的运动趋势仍然可以通过弹性对齐匹配上；如果轨迹方向、转弯形态和速度变化完全不一致，DTW 距离就会很大，连接分数应降低。")

    add_heading(doc, "9.6 无覆盖区域：时空可达性与路径推断", 2)
    add_para(doc, "当摄像头之间没有重复覆盖区域时，系统看不到车辆在中间道路上的真实运动，只能根据“从上游消失”和“在下游出现”之间的证据做推断。此时轨迹拼接不是几何轨迹对齐，而是候选边打分问题：两个 tracklet 是否可能属于同一目标。")
    add_table(
        doc,
        ["证据", "怎么计算", "为什么需要"],
        [
            ("ReID 外观", "计算两个 tracklet 的 avg_reid_vector 余弦相似度。", "无覆盖区域没有连续轨迹，外观是最重要的身份线索。"),
            ("时间可达性", "用下游出现时间减去上游消失时间，得到 travel_time。结合摄像头间距离估算速度。", "如果需要 2 分钟才能到达，但车辆 5 秒后就出现，基本不可信。"),
            ("路网拓扑", "根据摄像头连接关系、道路方向、可达路径判断上游是否能到下游。", "过滤物理上不连通或方向相反的摄像头组合。"),
            ("属性一致性", "比较颜色、车型、车牌等结构化属性。", "辅助 ReID，尤其在画质低或遮挡时提升稳定性。"),
            ("方向一致性", "比较离开上游和进入下游的运动方向。", "避免把相反车道或反方向运动误连。"),
            ("缺失惩罚", "如果缺少 ReID、缺少位置或时间不准，不直接否定，但降低置信度。", "让系统区分“证据不足”和“证据冲突”。"),
        ],
        [Inches(1.25), Inches(3.25), Inches(2.8)],
    )
    add_para(doc, "无覆盖区域的输出不应叫真实轨迹段，而应叫 inference segment。它表达的是：系统推断车辆可能从摄像头 A 到摄像头 B，中间没有可见画面，所以必须附带置信度、时间差、估计速度和主要证据。这样比直接画一条确定轨迹更严谨。")

    add_heading(doc, "9.7 观测链构建", 2)
    add_para(doc, "ObservationChainBuilder 把 Tracklet 看作图节点，把 CrossCameraEdge 看作有向边。以用户确认的锚点 Tracklet 为中心，向上游和下游搜索高置信连接，形成有序观测链。整体置信度使用连接边分数的几何平均，比简单算术平均更容易惩罚某个很低分的薄弱连接。")
    add_bullets(doc, [
        "上游扩展：查找时间更早且能连接到锚点的 tracklet。",
        "下游扩展：查找锚点之后最可能继续出现的 tracklet。",
        "环路控制：避免重复访问同一 tracklet 或同一摄像头。",
        "输出转换：把 tracklet 序列转换为观测节点、观测段、推断段和候选路径。",
        "可解释性：保留每条边的分项得分和 reasoning，便于后续人工复核。"
    ])

    add_heading(doc, "9.8 路径回溯结果怎么解释", 2)
    add_para(doc, "路径回溯不是把所有摄像头中的点机械连成一条线，而是围绕确认目标生成一条带证据的时空链。链中有两类内容：一种是系统真实看见的 observation，另一种是系统根据时空和身份证据推断出来的 inference。面试时一定要把这两类分开讲，否则容易被认为在伪造连续轨迹。")
    add_table(
        doc,
        ["结果类型", "含义", "来自哪里", "可信度解释"],
        [
            ("observation_node", "某个摄像头中真实观测到目标的节点。", "检测框、tracklet、关键帧、timestamp。", "置信度主要来自检测质量、ReID / CLIP 特征和属性一致性。"),
            ("observation_segment", "同一摄像头视野内的真实轨迹片段。", "单摄跟踪或标注轨迹。", "这是有画面证据的轨迹段，可信度通常高于推断段。"),
            ("inference_segment", "两个摄像头之间没有画面覆盖的推断连接。", "跨镜边评分、时间可达性、拓扑关系、ReID 和属性证据。", "必须展示置信度和证据组成，不能当作真实拍到的轨迹。"),
            ("candidate_path", "由多个 observation 和 inference 组成的候选完整路径。", "图搜索或路径排序。", "整体置信度受最弱连接影响，某一段低分会拉低整条路径。"),
        ],
        [Inches(1.4), Inches(2.0), Inches(2.0), Inches(1.9)],
    )
    add_code_block(
        doc,
        "例子：\n"
        "c001 observation：车辆在 08:01:12-08:01:20 被真实拍到\n"
        "c001 -> c004 inference：中间没有画面，系统根据 ReID、时间差和路网推断车辆可能到达 c004\n"
        "c004 observation：车辆在 08:03:05-08:03:18 被真实拍到\n\n"
        "最终展示时应该把 c001 和 c004 的观测段画成实线，把 c001 -> c004 的推断段画成虚线，并显示推断置信度。"
    )

    add_heading(doc, "10．端到端训练目标与损失函数", 1)
    add_para(doc, "如果把本项目改成端到端算法框架，核心思路是：前端检测和跟踪提供目标实例，双塔模型学习文本－图像检索空间，ReID 分支学习跨摄像头身份空间，属性分支学习可解释属性，边评分分支学习两个 tracklet 是否属于同一目标，路径分支学习整条观测链是否正确。总损失不是单一 loss，而是多任务加权损失。")
    add_code_block(
        doc,
        "total_loss = det_weight * detection_loss\n"
        "           + clip_weight * image_text_matching_loss\n"
        "           + reid_weight * identity_embedding_loss\n"
        "           + attr_weight * attribute_loss\n"
        "           + edge_weight * cross_camera_edge_loss\n"
        "           + path_weight * trajectory_path_loss\n"
        "           + uncertainty_weight * confidence_calibration_loss"
    )
    add_para(doc, "上面这行公式的意思是：端到端训练不是只优化一个目标，而是同时训练“能不能检测到目标”“文字和图片能不能对上”“跨镜是不是同一辆车”“属性是否识别正确”“两个轨迹片段能不能连接”“整条路径是否合理”“置信度是否可信”。每个 weight 是权重，用来控制某一类任务对总训练目标的影响。")
    add_table(
        doc,
        ["损失项", "监督信号", "优化目标", "面试讲法"],
        [
            ("L_det", "bbox、类别标签。", "让检测分支稳定输出车辆、行人、非机动车位置和类别。", "负责“看见目标”。"),
            ("L_clip", "文本－图像正负样本对。", "让匹配文本和图像在共享空间中靠近，不匹配样本远离。", "负责“文字找目标”。"),
            ("L_reid", "车辆 / 行人身份 ID。", "让同一身份跨摄像头 embedding 靠近，不同身份远离。", "负责“跨镜是不是同一个目标”。"),
            ("L_attr", "颜色、车型、背包、服饰等属性标签。", "提升细粒度属性识别，辅助检索和跨镜解释。", "负责“目标有什么可解释属性”。"),
            ("L_edge", "tracklet_i 与 tracklet_j 是否同一身份。", "学习跨镜候选边连接概率。", "负责“两个片段能不能连”。"),
            ("L_path", "真实轨迹链或正负路径。", "让正确观测链得分高于错误链。", "负责“整条轨迹是否合理”。"),
            ("L_uncertainty", "连接正确性、时间间隔、路径缺失等弱监督。", "校准推断段置信度，避免过度自信。", "负责“结果可信度是否可靠”。"),
        ],
        [Inches(1.0), Inches(1.45), Inches(2.5), Inches(2.35)],
    )
    add_table(
        doc,
        ["英文缩写", "中文含义", "一句话解释"],
        [
            ("loss", "损失函数", "模型预测错了要付出的代价，越小代表训练目标完成得越好。"),
            ("weight", "权重", "控制某个任务在总损失中的重要程度。"),
            ("embedding", "向量表征", "把文本、图像或轨迹片段变成一串数字，用于计算相似度。"),
            ("BCE", "二分类交叉熵", "用于训练“是 / 否”问题，例如两个 tracklet 是否同一车辆。"),
            ("Cross Entropy", "多分类交叉熵", "用于训练颜色、车型、身份 ID 等多类别分类问题。"),
            ("Triplet Loss", "三元组损失", "让同一车辆的向量更近，不同车辆的向量更远。"),
            ("margin", "间隔", "要求正样本分数至少比负样本高出一定距离。"),
        ],
        [Inches(1.45), Inches(1.55), Inches(4.3)],
    )

    add_heading(doc, "10.1 图文双塔对比损失", 2)
    add_para(doc, "Chinese-CLIP 分支用对比学习损失。一个 batch 中包含 N 对文本和图像，正确配对为正样本，其余 N-1 个图像或文本为负样本。")
    add_code_block(
        doc,
        "similarity_ij = cosine(text_embedding_i, image_embedding_j) / temperature\n\n"
        "text_to_image_loss: for each text, make the correct image rank first\n"
        "image_to_text_loss: for each image, make the correct text rank first\n\n"
        "image_text_matching_loss = 0.5 * (text_to_image_loss + image_to_text_loss)"
    )
    add_para(doc, "这个损失的直观含义是：如果文本是“蓝色轿车”，模型应该把它和真正的蓝色轿车图片拉近，同时把黑色轿车、蓝色 SUV、白色货车等不匹配图片推远。temperature 是温度系数，用来控制相似度分布的尖锐程度；temperature 越小，模型越强调最相似的样本。")
    add_para(doc, "训练时不能只随机采负样本。更有价值的是 hard negative，例如颜色相同但车型不同、车型相同但颜色不同、外观相似但不是同一车辆的样本。这样模型才会学习细粒度差异，而不是只学会很粗的颜色或大类别。")

    add_heading(doc, "10.2 ReID 身份损失", 2)
    add_para(doc, "ReID 分支建议使用 ID 分类损失和 Triplet Loss 组合。分类损失提供全局身份判别能力，Triplet Loss 直接优化 embedding 空间中的同车距离和异车距离。")
    add_code_block(
        doc,
        "identity_embedding_loss = id_classification_loss + triplet_weight * triplet_loss\n\n"
        "id_classification_loss: predict which vehicle ID this crop belongs to\n"
        "triplet_loss: distance(anchor, positive) should be smaller than distance(anchor, negative)"
    )
    add_para(doc, "这里的 anchor 是当前车辆图像，positive 是同一辆车的另一张图，negative 是另一辆车的图。Triplet Loss 要求模型满足一个很直观的关系：当前车和同一辆车的距离，要小于当前车和另一辆车的距离，并且最好小出一个 margin。margin 可以理解为安全间隔，防止模型只是勉强分开正负样本。")
    add_table(
        doc,
        ["样本", "含义", "例子"],
        [
            ("anchor", "当前目标图像。", "c001 中 V0034 的一帧车辆 crop。"),
            ("positive", "同一身份的另一张图。", "c004 中 V0034 的车辆 crop。"),
            ("negative", "不同身份的图。", "c004 中 V0088 或 V0102 的车辆 crop。"),
        ],
        [Inches(1.1), Inches(2.6), Inches(3.6)],
    )

    add_heading(doc, "10.3 属性和检测损失", 2)
    add_para(doc, "属性分支用于颜色、车型、背包、服饰等可解释字段。车辆颜色、车型通常用交叉熵；若类别不平衡，例如某些颜色样本很少，可以使用 Focal Loss。检测分支如果训练 YOLO 类模型，通常包括分类损失、框回归损失和 objectness 损失。")
    add_code_block(
        doc,
        "L_attr = L_color_ce + L_type_ce + L_bag_bce + L_clothing_ce\n\n"
        "Focal Loss: put more training focus on hard or rare samples\n\n"
        "L_det = L_cls + L_box + L_obj"
    )
    add_para(doc, "属性损失的作用是让模型学会可解释标签，例如颜色、车型、背包、服饰。检测损失的作用是让模型先把目标框出来。对于面试表达，可以不展开 YOLO 内部细节，只要说明检测损失通常由类别分类、框位置回归和目标存在性三部分组成。")

    add_heading(doc, "10.4 跨镜候选边损失", 2)
    add_para(doc, "跨镜拼接可以训练一个边评分网络。每条候选边由两个 tracklet 构成，输入包括 ReID 相似度、CLIP 相似度、属性一致性、车牌一致性、时间差、距离、方向差和拓扑可达性。标签 y_ij 表示两个 tracklet 是否属于同一目标。")
    add_code_block(
        doc,
        "edge_feature_ij = [reid_score, clip_score, attribute_score, plate_score, time_gap, distance, direction, topology]\n"
        "p_ij = sigmoid(MLP(edge_feature_ij))\n\n"
        "cross_camera_edge_loss = BCE(edge_label, predicted_edge_probability)"
    )
    add_para(doc, "edge_label 是真实标签：如果两个 tracklet 属于同一辆车，edge_label 为 1；如果不是同一辆车，edge_label 为 0。predicted_edge_probability 是模型输出的连接概率。BCE 会惩罚两类错误：真实同车却预测为不能连接，或者真实不同车却预测为可以连接。")
    add_para(doc, "除了 BCE，还可以使用排序损失，让真实连接边得分高于错误连接边。")
    add_code_block(
        doc,
        "edge_ranking_loss = max(0, margin - positive_edge_score + negative_edge_score)"
    )

    add_heading(doc, "10.5 轨迹路径损失与解码", 2)
    add_para(doc, "轨迹回溯的最终目标不是单条边，而是一整条观测链。因此可以把候选路径作为训练样本，让真实路径得分高于错误路径。路径得分可以是边得分的加权和、几何平均，或由序列模型输出。")
    add_code_block(
        doc,
        "path_score = sum(edge_scores) - gap_weight * path_gap_penalty - missing_weight * missing_observation_penalty\n\n"
        "trajectory_path_loss = max(0, margin - true_path_score + false_path_score)"
    )
    add_para(doc, "路径损失的目标是让整条正确路径比错误路径得分更高。比如真实路径是 c001 -> c004 -> c009，而错误路径是 c001 -> c006 -> c009，那么模型不仅要判断单条边是否合理，还要判断整条路径的时间顺序、空间可达性和证据连续性是否更可信。")
    add_callout(
        doc,
        "端到端面试表达",
        "可以这样讲：我不是把 CLIP、ReID 和规则简单串起来，而是把系统拆成可监督的多任务目标。文本检索由 L_clip 优化，身份一致性由 L_reid 优化，属性解释由 L_attr 优化，跨镜连接由 L_edge 优化，最终轨迹链由 L_path 优化。工程部署时可离线建索引，但训练目标是端到端联合优化。",
    )

    add_heading(doc, "11．算法亮点与创新点", 1)
    add_para(doc, "这个项目的算法亮点不在于单独使用某个模型，而在于把跨模态检索、ReID 身份判别和交通时空约束组织成一个可落地的研判闭环。面试时建议把亮点讲成“为什么这样设计能解决交通视频里的真实问题”。")
    add_table(
        doc,
        ["亮点", "具体做法", "解决的问题", "面试表达"],
        [
            ("图文双塔快速检索", "图像塔离线编码图库，文本塔在线编码查询，FAISS / Qdrant 做近邻检索。", "避免每次查询都遍历视频或逐对跑大模型。", "双塔结构带来可预计算和可扩展检索。"),
            ("属性粗筛 + 语义精排", "先用颜色、车型、车牌等结构化属性缩小候选，再用 Chinese-CLIP 排序。", "减少纯向量检索的细粒度属性误召回。", "结构化规则保证可控，CLIP 保证语义泛化。"),
            ("ReID + 交通约束融合", "跨镜边评分同时考虑 ReID 外观、属性、车牌、时间、空间、方向。", "单纯 ReID 易受光照、视角和遮挡影响。", "把视觉特征和交通领域知识融合，提升跨镜关联可靠性。"),
            ("离散观测链建模", "把摄像头内轨迹定义为观测段，把摄像头间连接定义为推断段。", "道路摄像头覆盖不连续，不能伪造成连续轨迹。", "这是比“画一条线”更严谨的轨迹表达。"),
            ("不确定性可视化", "输出每段连接置信度、推断旅行时间、实际时间差和证据分项。", "让研判人员知道哪些部分是强证据，哪些部分是推断。", "算法结果不只给结论，还给可信度和证据。"),
            ("人机确认闭环", "检索阶段返回候选，用户确认后再进行轨迹回溯。", "避免初始文本误召回直接污染整条轨迹。", "确认环节把任务从模糊检索转为实例级追踪。"),
        ],
        [Inches(1.2), Inches(2.2), Inches(1.95), Inches(1.95)],
    )
    add_heading(doc, "11.1 可以重点强调的创新点", 2)
    add_bullets(doc, [
        "从“图片级检索”升级到“目标级时空回溯”：系统不止返回相似图，还输出目标经过哪些摄像头、什么时间出现、证据来自哪里。",
        "从“单模态视觉匹配”升级到“多证据融合评分”：ReID、车牌、属性、时间可达性、空间拓扑和方向一致性共同决定跨镜连接。",
        "从“连续轨迹假设”升级到“离散观测链”：明确区分真实观测和推断连接，更符合城市道路监控的实际覆盖条件。",
        "从“黑盒结果”升级到“可解释研判”：候选得分和跨镜边得分都有分项，方便面试中讲消融实验和错误分析。"
    ])

    add_heading(doc, "12．接口、前端与部署架构", 1)
    add_table(
        doc,
        ["接口", "方法", "作用", "输入输出重点"],
        [
            ("/api/v1/search/query", "POST", "文本查询检索。", "输入 query_text、top_k；输出 candidates、分数和关键帧。"),
            ("/api/v1/search/plate", "POST", "车牌精确查询。", "输入 plate_number；按车牌字段匹配。"),
            ("/api/v1/confirm/target", "POST", "目标确认。", "把候选目标转为后续回溯锚点。"),
            ("/api/v1/backtrack/trace", "POST", "锚点回溯。", "输出 observation_nodes、segments、candidate_paths。"),
            ("/api/v1/backtrack/trajectory", "POST", "完整跨镜时间线。", "按 instance_id 或 track_id 解析 vehicle_id 后返回多摄像头轨迹。"),
            ("/api/v1/dashboard/stats", "GET", "统计看板。", "输出检测数、轨迹数、摄像头等统计信息。"),
        ],
        [Inches(1.85), Inches(0.65), Inches(1.65), Inches(3.15)],
    )
    add_para(doc, "部署上，docker-compose.yml 编排 backend、frontend 和 qdrant 三个服务。backend 使用 FastAPI + Uvicorn 并预留 NVIDIA GPU；frontend 使用 Streamlit；qdrant 用于向量库持久化。后端端口 8000，前端端口 8501，Qdrant 端口 6333/6334。")
    add_code_block(
        doc,
        "backend: FastAPI + CUDA，挂载 /data、/models、/output、/logs\n"
        "frontend: Streamlit，通过 API_BASE_URL=http://backend:8000 调用后端\n"
        "qdrant: 向量数据库，持久化到 docker_data/qdrant\n"
        "network: traffic-net bridge"
    )

    add_heading(doc, "13．面试可讲重点与追问回答", 1)
    add_table(
        doc,
        ["面试问题", "推荐回答"],
        [
            ("这个项目的核心难点是什么？", "难点不是单个模型，而是把多摄像头视频转成可检索、可确认、可回溯的数据闭环。检索要兼顾语义和结构化属性，回溯要承认摄像头覆盖不连续，用观测段和推断段区分证据强弱。"),
            ("为什么需要用户确认？", "自然语言检索本质是候选召回，可能存在相似但不是同一目标的结果。用户确认后，系统可以从文本级检索切换为实例级回溯，后续围绕 instance_id、track_id 或 vehicle_id 展开，误差传播更可控。"),
            ("CLIP 在项目里解决什么？", "Chinese-CLIP 是图文双塔模型，解决中文描述与目标图像之间的跨模态匹配，尤其是用户不按字段输入时的语义召回。但 CLIP 不是身份识别模型，所以要结合属性、车牌、ReID 和时空约束。"),
            ("ReID 在项目里解决什么？", "ReID 解决跨摄像头外观身份一致性问题。它不理解文本，但能判断两个不同摄像头里的目标是否外观相似。项目中跨镜评分优先使用 ReID 相似度，缺失时才回退到 CLIP 图像特征。"),
            ("为什么说是端到端？", "端到端不是指部署时每次都从原始视频重新跑全链路，而是指训练目标覆盖从文本－图像检索、身份表征、属性识别、跨镜边连接到最终路径解码的全过程。工程上可以离线建索引，算法上通过 L_total 联合优化。"),
            ("损失函数怎么设计？", "总损失由 L_det、L_clip、L_reid、L_attr、L_edge、L_path 组成。CLIP 用双塔对比损失，ReID 用 ID CE + Triplet，属性用 CE / Focal，跨镜边用 BCE 或 ranking loss，轨迹链用 path ranking loss。"),
            ("跨镜拼接为什么不能只靠 ReID？", "道路监控存在视角、光照、遮挡和分辨率变化，ReID 容易误匹配。交通场景有天然约束，如时间可达性、道路拓扑、方向和车牌，这些约束能显著降低不合理连接。"),
            ("如何评价轨迹可信度？", "单摄内观测是强证据，摄像头间连接是推断证据。连接置信度由外观、属性、车牌、时间、空间、方向和惩罚项共同决定；整条链可用连接分数的几何平均聚合。"),
            ("项目中你最能体现算法能力的点？", "我会讲双塔图文检索、ReID 跨镜外观匹配和跨镜多证据评分。双塔检索体现效率与语义匹配的平衡；ReID 和交通约束融合体现把视觉模型输出转成可解释轨迹证据。"),
            ("当前实现有什么边界？", "工程原型中已有离线预处理、检索、按 vehicle_id 聚合和轨迹展示；端到端版本需要进一步补齐统一训练数据、跨镜候选图监督、ReID 训练权重和 path-level 损失训练。"),
        ],
        [Inches(1.9), Inches(5.4)],
    )

    add_heading(doc, "14．当前边界与后续优化方向", 1)
    add_para(doc, "为了面试表达准确，需要区分“工程实现阶段”和“端到端算法目标”。当前项目已经具备预处理、检索、候选返回、按 vehicle_id 聚合轨迹和前端展示闭环；端到端版本需要把这些模块进一步统一到多任务训练框架中，补齐训练样本构造、边监督、路径监督和端到端评测。")
    add_table(
        doc,
        ["方向", "当前状态", "优化建议"],
        [
            ("检索精度", "工程上已有属性粗筛 + Chinese-CLIP 融合排序。", "建立人工标注 query－target 测试集，并用 L_clip 微调双塔模型，评估 Recall@K、MRR、P@K。"),
            ("双塔模型", "当前可使用预训练 Chinese-CLIP 做零样本或弱监督检索。", "用交通目标图文对做端到端对比学习微调，加入 hard negative，例如“蓝色轿车”和“蓝色 SUV”。"),
            ("ReID 特征", "FeatureExtractor 和评分器支持 ReID，跨镜评分优先使用 ReID 外观相似度。", "接入车辆 ReID / FastReID / OSNet 训练权重，使用 L_id_ce + L_triplet 训练身份 embedding。"),
            ("向量索引", "FAISS 文件索引和 Qdrant 存储均有代码路径。", "统一线上索引方案，保存 det_id / track_id 到向量行号的严格映射，并支持增量更新。"),
            ("跨镜拼接", "评分器和观测链构建器已模块化；接口层可按 vehicle_id 聚合。", "把候选边构造成监督样本，训练 L_edge；把真实轨迹链构造成路径样本，训练 L_path。"),
            ("拓扑约束", "配置中已有摄像头元数据、方向和场景信息。", "接入真实道路拓扑、路段长度、限速和转向限制，改进时间可达性和空间可达性分数。"),
            ("不确定性", "输出观测段、推断段和置信度。", "把时间间隔、路径分叉、中间摄像头缺失和 ReID 低相似度转成风险等级，并在前端可视化。"),
            ("工程化", "Docker Compose 已编排后端、前端、Qdrant。", "增加模型预热、索引版本管理、请求日志、异常告警、GPU 内存监控和批量离线任务调度。"),
        ],
        [Inches(1.2), Inches(2.55), Inches(3.55)],
    )
    add_heading(doc, "14.1 后续优化路线图", 2)
    add_table(
        doc,
        ["阶段", "目标", "具体动作", "验收指标"],
        [
            ("短期", "把端到端训练数据构造出来。", "构建中文 query－target 配对、同 ID / 异 ID ReID 样本、tracklet 正负边和真实轨迹链。", "训练样本覆盖率、正负样本比例、错误案例清单。"),
            ("中期", "训练多任务表征。", "联合或分阶段训练 L_clip、L_reid、L_attr，得到统一目标表征。", "Recall@K、mAP、属性准确率、ReID mAP。"),
            ("中期", "训练跨镜边和路径。", "以 tracklet 为节点生成候选图，用 L_edge 学习边概率，用 L_path 学习路径排序。", "边 AUC、轨迹链 Top-1 / Top-3 命中率、IDF1。"),
            ("长期", "从端到端算法升级为可持续平台。", "加入增量索引、在线日志、模型版本管理、数据漂移监控和人工反馈闭环。", "可复现实验、可回滚模型、可解释错误分析。"),
        ],
        [Inches(0.8), Inches(1.45), Inches(3.15), Inches(1.9)],
    )

    add_heading(doc, "附录：一分钟项目介绍模板", 1)
    add_para(
        doc,
        "我做的是一个面向道路监控视频的端到端交通目标检索与跨镜轨迹回溯平台。整体不是简单串联 CLIP、ReID 和规则，而是把多摄像头视频、中文文本查询、目标检测、身份表征、属性识别、跨镜候选边和最终轨迹链放进统一多任务框架。模型前端从视频中得到目标实例和单摄 tracklet，中间分支同时学习 Chinese-CLIP 双塔图文对齐、ReID 身份 embedding 和颜色车型等属性，后端把 tracklet 构成候选图，用边评分网络判断两个片段是否同一目标，再解码成摄像头经过序列、观测段和推断段。训练上设计总损失 L_total = L_det + L_clip + L_reid + L_attr + L_edge + L_path，其中 CLIP 用双塔对比损失，ReID 用 ID 分类加 Triplet Loss，跨镜边用 BCE 或 ranking loss，轨迹链用 path ranking loss。部署时可以离线建索引、在线检索，但算法目标是端到端联合优化从“文字找目标”到“跨镜轨迹回溯”的完整链路。"
    )

    # Footer
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("交通目标检索与跨镜轨迹回溯平台技术说明书")
        set_run_font(run, size=8.5, color=MUTED)

    doc.save(OUT)
    patch_font_ooxml(OUT)
    return OUT


def patch_font_ooxml(path: Path):
    """Ensure all runs carry SimSun eastAsia and Times New Roman ascii/hAnsi."""
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                text = data.decode("utf-8")
                text = text.replace('w:ascii="Calibri"', 'w:ascii="Times New Roman"')
                text = text.replace('w:hAnsi="Calibri"', 'w:hAnsi="Times New Roman"')
                text = text.replace('w:eastAsia="Calibri"', 'w:eastAsia="SimSun"')
                text = text.replace('w:eastAsiaTheme="minorEastAsia"', 'w:eastAsia="SimSun"')
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


if __name__ == "__main__":
    out = build_document()
    print(out)
